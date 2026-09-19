"""Carga en la gestoría las facturas de una carpeta EMITIDAS/RECIBIDAS de un cliente.

Hecho para la contabilidad de Modernia Home & Investment (2026-09-19): facturas hechas
en Excel (hoja con "Nº Factura", "CLIENTE", "DESCRIPCIÓN" y "BASE / IVA / TOTAL
FACTURA") con su PDF al lado, y autofacturas de Fuxiona solo en PDF.

Para cada factura: pide a producción una subida firmada (/api/ingest_facturas_presign),
sube el PDF a S3 y la da de alta (/api/ingest_facturas_ocr) con los datos leídos del
Excel, que mandan sobre el OCR, y con el cliente cuya contabilidad se lleva. El alta es
idempotente: una factura ya cargada (misma clave de duplicado) no se repite.

Sin --aplicar solo enseña lo que haría. La clave de ingesta se lee de APP_INGEST_API_KEY
(en el .env); el script no la imprime.

  python scripts/cargar_facturas_de_carpeta.py CARPETA --empresa-id ID --cliente-id ID [--aplicar]
"""

import argparse
import datetime
import json
import os
import re
import sys
import urllib.request
from pathlib import Path

BASE_URL = os.environ.get("CRM_BASE_URL", "https://crm.verifika2.com")


def _r2(v):
    return round(float(v or 0) + 1e-9, 2)


def _fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(v or ""))
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None


def leer_excel(ruta):
    """Datos de una factura hecha en Excel: la hoja que tiene 'Nº Factura' y 'TOTAL FACTURA'."""
    from openpyxl import load_workbook

    wb = load_workbook(ruta, data_only=True, read_only=True)
    for ws in wb.worksheets:
        filas = [[v for v in r if v not in (None, "")] for r in ws.iter_rows(max_row=80, values_only=True)]
        filas = [f for f in filas if f]
        datos = {}
        for i, f in enumerate(filas):
            cabeza = str(f[0]).strip().upper()
            if cabeza == "FECHA" and len(f) > 1:
                datos["fecha"] = _fecha(f[1])
                if len(f) > 2 and str(f[2]).strip().upper() == "CLIENTE" and i + 2 < len(filas):
                    datos["tercero"] = str(filas[i + 1][0]).strip()
                    datos["nif"] = re.sub(r"[\s.\-]", "", str(filas[i + 2][0])).upper()
            if cabeza.startswith("Nº FACTURA") and len(f) > 1:
                datos["numero"] = str(f[1]).strip()
            if cabeza.startswith("DESCRIPCI") and i + 1 < len(filas):
                datos["descripcion"] = str(filas[i + 1][0]).strip()
            if [str(x).strip().upper() for x in f[:3]] == ["BASE", "IVA", "TOTAL FACTURA"] and i + 1 < len(filas):
                base, iva, _total = filas[i + 1][:3]
                datos["base_imponible"] = _r2(base)
                datos["cuota_iva"] = _r2(iva)
                # El total, como suma de lo redondeado: el Excel arrastra decimales (7.999,9997).
                datos["total"] = _r2(datos["base_imponible"] + datos["cuota_iva"])
        # Sin descripción en el Excel (la 17/2025 traía solo "1900"), la del nombre del
        # fichero: "17 Estudio Ciudad Jardin Sur - Cuota asociacion agosto".
        if not re.search(r"[A-Za-zÁÉÍÓÚáéíóúñÑ]", str(datos.get("descripcion") or "")):
            datos["descripcion"] = re.sub(r"^\d+\s*", "", Path(ruta).stem).strip()
        if all(datos.get(k) for k in ("numero", "fecha", "total")):
            if datos["base_imponible"]:
                datos["iva_pct"] = round(datos["cuota_iva"] * 100 / datos["base_imponible"])
            return datos
    return None


def leer_autofactura_fuxiona(ruta):
    """Autofactura de Fuxiona (solo PDF): 'Factura: MHI07/25', 'Fecha: dd/mm/aaaa', totales."""
    from pypdf import PdfReader

    texto = "\n".join(p.extract_text() or "" for p in PdfReader(str(ruta)).pages)
    num = re.search(r"Factura:\s*(\S+)", texto)
    fecha = re.search(r"Fecha:\s*(\d{1,2}/\d{1,2}/\d{4})", texto)
    cliente = re.search(r"Datos del cliente\s*\n(.+)\n([A-Z0-9]{9})", texto)
    tot = re.search(r"TOTAL FACTURA\s*\n\s*([\d.,]+) €\s+([\d.,]+)\s+([\d.,]+) €\s+[\d.,]+\s+([\d.,]+) €\s+([\d.,]+) €", texto)
    if not (num and fecha and tot):
        return None
    eu = lambda s: float(s.replace(".", "").replace(",", "."))  # noqa: E731
    base, iva = eu(tot.group(1)), eu(tot.group(3))
    return {
        "numero": num.group(1), "fecha": _fecha(fecha.group(1)),
        "tercero": cliente.group(1).strip() if cliente else "", "nif": cliente.group(2) if cliente else "",
        "descripcion": "Captación (autofactura Fuxiona)", "base_imponible": _r2(base), "cuota_iva": _r2(iva),
        "total": _r2(base + iva), "iva_pct": round(eu(tot.group(2))), "irpf": eu(tot.group(4)),
    }


def _texto_documento(ruta):
    """Texto de un PDF o de un Word (.docx: las emitidas 16–24/2026 de GAPP solo están así)."""
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".docx":
        import zipfile

        with zipfile.ZipFile(ruta) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        xml = re.sub(r"</w:p>", "\n", xml)
        xml = re.sub(r"<w:tab/>", " ", xml)
        return re.sub(r"<[^>]+>", "", xml)
    from pypdf import PdfReader

    return "\n".join(p.extract_text() or "" for p in PdfReader(str(ruta)).pages)


def leer_factura_gapp(ruta):
    """Factura emitida por GAPP (Word o su PDF): 'FACTURA Nº 01_2026', 'Cliente : X', 'CIF', totales."""
    texto = _texto_documento(ruta)
    plano = re.sub(r"[ \t]+", " ", texto)
    num = re.search(r"FACTURA\s+N[ºo°]\s*(\d+)_(\d{4})", plano)
    fecha = re.search(r"Fecha:?[\s|]*(\d{1,2}/\d{1,2}/\d{4})", plano.replace("\n", " | "))
    cliente = re.search(r"Cliente\s*:\s*(.+?)(?:\s+CIF:|\n)", plano)
    cif = re.search(r"CIF:\s*([A-Z]\d{8}|[A-Z]\d{7}[A-Z0-9])", plano)
    base = re.search(r"Base\s+imponible\s+([\d.]+,\d{2})", plano)
    iva = re.search(r"IVA\s+21\s*%\s+([\d.]+,\d{2})", plano)
    total = re.search(r"Total\s+([\d.]+,\d{2})", plano)
    desc = re.search(r"DESCRIPCI[ÓO]N\s+IMPORTE\s*\n(.+)", plano)
    if not (num and fecha and base and iva):
        return None
    eu = lambda s: float(s.replace(".", "").replace(",", "."))  # noqa: E731
    b, i = eu(base.group(1)), eu(iva.group(1))
    aviso = ""
    if total and abs(eu(total.group(1)) - (b + i)) > 0.005:
        # La factura se contradice: manda la cifra que cuadra con el 21 % de la base.
        # 64/2026: IVA impreso 695,13 (mal) y total 4.008,73 (bien) → se toma el total.
        # 32/2026: IVA 100,80 (bien) y total 580,00 (mal) → se toma el IVA.
        iva_21 = round(b * 0.21, 2)
        t = eu(total.group(1))
        if abs(i - iva_21) > 0.01 and abs((t - b) - iva_21) <= 0.01:
            aviso = f"el IVA impreso ({i:.2f}) no es el 21 % de la base; se toma el total impreso {t:.2f}"
            i = t - b
        else:
            aviso = f"el total impreso ({t:.2f}) no es base + IVA ({b + i:.2f}); se toma base + IVA"
    return {
        "aviso": aviso,
        "numero": f"{int(num.group(1)):02d}/{num.group(2)}", "fecha": _fecha(fecha.group(1)),
        "tercero": re.sub(r"\s+", " ", cliente.group(1)).strip() if cliente else "",
        "nif": cif.group(1) if cif else "",
        "descripcion": re.sub(r"\s+", " ", desc.group(1)).strip()[:200] if desc else "Trabajos de ascensores",
        "base_imponible": _r2(b), "cuota_iva": _r2(i), "total": _r2(b + i), "iva_pct": 21,
    }


LECTORES_PDF = (leer_autofactura_fuxiona, leer_factura_gapp)


def _leer_pdf(pdf):
    for lector in LECTORES_PDF:
        try:
            datos = lector(pdf)
        except Exception:
            datos = None
        if datos:
            return datos
    return None


def facturas_de_la_carpeta(carpeta, tipo):
    """[(pdf, datos, origen)] de EMITIDAS o RECIBIDAS; las autofacturas repetidas, fuera."""
    raiz = Path(carpeta) / tipo
    salida, avisos = [], []
    for xlsx in sorted(raiz.rglob("*.xlsx")):
        datos = leer_excel(xlsx)
        pdf = xlsx.with_suffix(".pdf")
        if not pdf.exists():
            candidatos = sorted(xlsx.parent.glob(xlsx.stem.split(" ")[0] + " *.pdf"))
            pdf = candidatos[0] if candidatos else None
        if not datos:
            avisos.append(f"No se pudo leer {xlsx.name}")
            continue
        salida.append((pdf, datos, xlsx.name))
    con_excel = {(d["fecha"], d["total"]) for _p, d, _o in salida}
    documentos = sorted([*raiz.rglob("*.pdf"), *raiz.rglob("*.PDF"), *raiz.rglob("*.docx")], key=lambda x: str(x))
    for pdf in dict.fromkeys(documentos):
        if any(p == pdf for p, _d, _o in salida) or pdf.with_suffix(".xlsx").exists():
            continue
        datos = _leer_pdf(pdf)
        if not datos:
            avisos.append(f"PDF sin Excel que no sé leer: {pdf.name}")
            continue
        if (datos["fecha"], datos["total"]) in con_excel:
            avisos.append(f"{pdf.name} ({datos['numero']}) es la misma operación que una factura del Excel: no se carga dos veces")
            continue
        repetida = next((o for _p, d, o in salida if d["numero"] == datos["numero"]), None)
        if repetida:
            avisos.append(f"{pdf.name}: el nº {datos['numero']} ya está en {repetida}; no se carga dos veces")
            continue
        if datos.get("aviso"):
            avisos.append(f"{pdf.name}: {datos['aviso']}")
        salida.append((pdf, datos, pdf.name))
    return salida, avisos


def _post(ruta, cuerpo, clave):
    req = urllib.request.Request(
        BASE_URL + ruta, data=json.dumps(cuerpo).encode(), method="POST",
        headers={"Content-Type": "application/json", "X-API-Key": clave},
    )
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read().decode())


def cargar(pdf, datos, *, empresa_id, cliente_id, tipo, clave):
    anio = (datos.get("fecha") or "")[:4]
    firma = _post("/api/ingest_facturas_presign", {
        "empresa_id": empresa_id, "tipo": tipo, "year": anio, "filename": pdf.name, "content_type": "application/pdf",
    }, clave)
    tipo_mime = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                 if pdf.suffix.lower() == ".docx" else "application/pdf")
    firma = firma if pdf.suffix.lower() != ".docx" else _post("/api/ingest_facturas_presign", {
        "empresa_id": empresa_id, "tipo": tipo, "year": anio, "filename": pdf.name, "content_type": tipo_mime,
    }, clave)
    subida = urllib.request.Request(firma["url"], data=pdf.read_bytes(), method="PUT",
                                    headers={"Content-Type": tipo_mime})
    with urllib.request.urlopen(subida, timeout=120) as r:
        if r.status not in (200, 204):
            raise RuntimeError(f"S3 respondió {r.status}")
    campos = {k: datos[k] for k in ("numero", "fecha", "nif", "tercero", "descripcion", "base_imponible",
                                    "cuota_iva", "total", "iva_pct") if datos.get(k) not in (None, "")}
    return _post("/api/ingest_facturas_ocr", {
        "empresa_id": empresa_id, "s3_key": firma["key"], "filename": pdf.name, "tipo": tipo,
        "cliente_id": cliente_id, **campos,
    }, clave)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("carpeta")
    ap.add_argument("--empresa-id", required=True)
    ap.add_argument("--cliente-id", required=True)
    ap.add_argument("--tipo", choices=("EMITIDAS", "RECIBIDAS"), default="EMITIDAS")
    ap.add_argument("--aplicar", action="store_true")
    a = ap.parse_args()
    facturas, avisos = facturas_de_la_carpeta(a.carpeta, a.tipo)
    for aviso in avisos:
        print("AVISO:", aviso)
    total = 0.0
    for pdf, d, origen in facturas:
        total += d["total"]
        print(f"{d['fecha']}  {d['numero']:<10} {d.get('tercero', '')[:34]:<34} base {d['base_imponible']:>10.2f}  "
              f"IVA {d['cuota_iva']:>8.2f}  total {d['total']:>10.2f}  {'PDF' if pdf else 'SIN PDF'}  ← {origen}")
    print(f"{len(facturas)} facturas · total {total:.2f}")
    if not a.aplicar:
        print("(prueba: no se ha cargado nada; añade --aplicar)")
        return
    clave = os.environ.get("APP_INGEST_API_KEY", "").strip()
    if not clave:
        sys.exit("Falta APP_INGEST_API_KEY en el entorno (.env).")
    fallos = 0
    for pdf, d, origen in facturas:
        if not pdf:
            print("SIN PDF, no se carga:", origen)
            fallos += 1
            continue
        try:
            r = cargar(pdf, d, empresa_id=a.empresa_id, cliente_id=a.cliente_id, tipo=a.tipo, clave=clave)
            print("OK" if r.get("ok") else "??", d["numero"], "duplicada" if r.get("duplicate") else "", r.get("factura_id") or r.get("id") or "")
        except Exception as exc:  # seguir con las demás y contarlo
            fallos += 1
            detalle = getattr(exc, "read", lambda: b"")().decode(errors="replace")[:300]
            print("ERROR", d["numero"], exc, detalle)
    print(f"Terminado: {len(facturas) - fallos} cargadas, {fallos} con error.")


if __name__ == "__main__":
    main()
